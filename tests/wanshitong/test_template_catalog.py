"""模板只检索目录项，不把占位符或示例正文作为事实。"""

from __future__ import annotations

import io

from docx import Document

from rag_app.wanshitong.template_catalog import (
    is_template_source,
    searchable_upload_content,
)


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("仅用于测试的占位正文，不应进入检索。")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_template_directory_or_filename_uses_catalog_only_docx() -> None:
    original = _docx_bytes()
    for path in (
        "04 开发中心/流程模板/设计说明书.docx",
        "04 开发中心/流程/设计说明书模板.docx",
    ):
        assert is_template_source(path)
        content = searchable_upload_content(
            original,
            source_relative_path=path,
            document_title="设计说明书",
        )
        catalog = Document(io.BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in catalog.paragraphs)
        assert "模板目录项：设计说明书" in text
        assert "参考原始模板" in text
        assert "占位正文" not in text


def test_non_template_upload_keeps_original_bytes() -> None:
    original = _docx_bytes()
    assert not is_template_source("04 开发中心/工作指引/设计规范.docx")
    assert (
        searchable_upload_content(
            original,
            source_relative_path="04 开发中心/工作指引/设计规范.docx",
            document_title="设计规范",
        )
        is original
    )
