"""模板解析保留真实填写指引与逐字来源。"""

from __future__ import annotations

import hashlib
import io

import pytest
from docx import Document

from rag_app.adapters.chunkers.weknora.chunker import WeKnoraChunkerAdapter
from rag_app.adapters.chunkers.weknora.reading_domain import (
    plan_reading_domains,
)
from rag_app.adapters.parsers.docx import DocxOoxmlV4Parser
from rag_app.core.identifiers import deterministic_id
from rag_app.core.models import (
    ChunkingContext,
    DocumentRef,
    ParseContext,
    ParseSource,
    WeKnoraChunkingPolicy,
)
from rag_app.core.models.common import freeze_json_object
from rag_app.core.policies import ParsingPolicy
from rag_app.wanshitong.template_parser import WanshitongTemplateParser

_GUIDANCE = "本项目要做什么？解决什么业务问题或用户痛点？"
_EXAMPLE = "用 SMART 原则写出量化目标（例如：上线后支持一定并发用户）。"
_SCOPE = "明确包含功能（必做）、不包含功能（明确排除）。"


def _template_bytes(*, leading_paragraphs: int) -> bytes:
    document = Document()
    document.add_paragraph("XX 产品")
    for index in range(leading_paragraphs):
        document.add_paragraph(f"填写前说明 {index}")
    document.add_heading("项目基本信息", level=1)
    document.add_paragraph(_GUIDANCE)
    document.add_paragraph(_EXAMPLE)
    document.add_heading("项目范围", level=1)
    document.add_paragraph(_SCOPE)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    ("relative_path", "leading_paragraphs"),
    [
        ("业务/模板库/甲计划.docx", 0),
        ("业务/项目文件/乙计划模板.docx", 3),
    ],
)
def test_template_ir_and_reading_view_keep_guidance(
    relative_path: str,
    leading_paragraphs: int,
) -> None:
    content = _template_bytes(leading_paragraphs=leading_paragraphs)
    name = relative_path.rsplit("/", maxsplit=1)[-1]
    parser = WanshitongTemplateParser(DocxOoxmlV4Parser())
    result = parser.parse(
        ParseSource(
            media_type="application/octet-stream",
            display_name=name,
            extension=".docx",
            content=content,
        ),
        ParsingPolicy(),
        ParseContext(
            document=DocumentRef(
                project_id=f"prj_{'1' * 32}",
                knowledge_base_id=f"kb_{'2' * 32}",
                document_id=f"doc_{'3' * 32}",
                display_name=name,
                metadata=freeze_json_object(
                    {"source_relative_path": relative_path}
                ),
            )
        ),
    )

    assert parser.descriptor.name == "wanshitong-template-source-v3"
    assert result.report.parser_id == "docx-ooxml-v4"
    assert (
        result.document_ir.source.content_sha256
        == hashlib.sha256(content).hexdigest()
    )
    text_nodes = {
        node.text_payload.exact_text: node
        for node in result.document_ir.nodes
        if node.text_payload is not None
    }
    for text in (_GUIDANCE, _EXAMPLE, _SCOPE):
        assert text in text_nodes
        assert text_nodes[text].anchor.source_end_char == len(text)
    assert "XX 产品" in text_nodes
    assert not any("模板正文未入库" in text for text in text_nodes)

    domains = plan_reading_domains(result.document_ir, WeKnoraChunkingPolicy())
    reading_text = "\n".join(domain.view.text for domain in domains)
    assert all(text in reading_text for text in (_GUIDANCE, _EXAMPLE, _SCOPE))
    assert "例如" in reading_text
    assert "XX 产品" in reading_text

    chunker = WeKnoraChunkerAdapter()
    preview = chunker.preview(
        result.document_ir,
        ChunkingContext(
            chunker_fingerprint=chunker.fingerprint,
            index_revision_id=deterministic_id(
                "irev", "template-guidance", name
            ),
        ),
    )
    indexed_text = "\n".join(
        chunk.citation_text for chunk in preview.result.chunks
    )
    assert all(text in indexed_text for text in (_GUIDANCE, _EXAMPLE, _SCOPE))
    assert preview.result.report.source_span_coverage == 1.0
