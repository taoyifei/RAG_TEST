"""V3-00.6 DOCX 来源、空白结构和 token 边界回归。"""

from __future__ import annotations

import hashlib
import io
import unicodedata

import pytest
from docx import Document
from pydantic import ValidationError

from rag_app.adapters.chunkers.docx_structural.atoms import (
    AtomicUnit,
    SourceFragment,
)
from rag_app.adapters.chunkers.docx_structural.context import embedding_text
from rag_app.adapters.chunkers.docx_structural.rendering import (
    render_fragments,
    separator_fragment,
)
from rag_app.adapters.chunkers.docx_structural.splitting import split_atom
from rag_app.adapters.chunkers.docx_structural.validation import (
    validate_chunks,
)
from rag_app.adapters.tokenizers import DeterministicUtf8TokenCounter
from rag_app.core.identifiers import deterministic_id
from rag_app.core.models import (
    Chunk,
    ChunkingPolicy,
    ChunkRole,
    NodeKind,
    SourceAnchor,
    SourceSpanKind,
    StoryKind,
)
from tests.adapters.chunkers.test_docx_structural import _chunk
from tests.adapters.parsers.docx.fixtures import build_package, parse_package
from tests.fixtures.docx_v4.generate_fixtures import _list_item, _numbering


def _python_docx_with_whitespace_section() -> bytes:
    """用 python-docx 生成与手写 OOXML 独立的空白段落样例。"""
    document = Document()
    document.add_heading("可见章节", level=1)
    document.add_paragraph("设备 ZK-204 的巡检周期为九天。")
    document.add_heading("空白章节", level=1)
    document.add_paragraph(" \t\u00a0")
    document.add_heading("符号章节", level=1)
    document.add_paragraph("；")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    "content",
    (
        build_package(
            '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>'
            "<w:r><w:t>可见章节</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>设备 AX-17 的校验周期为十一天。</w:t>"
            "</w:r></w:p>"
            '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>'
            "<w:r><w:t>空白章节</w:t></w:r></w:p>"
            '<w:p><w:r><w:t xml:space="preserve">        </w:t>'
            "</w:r></w:p>"
            '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>'
            "<w:r><w:t>符号章节</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>。</w:t></w:r></w:p>"
        ),
        _python_docx_with_whitespace_section(),
    ),
    ids=("minimal-ooxml", "python-docx"),
)
def test_whitespace_only_sections_are_non_citable_but_punctuation_remains(
    content: bytes,
) -> None:
    """空白段落不构造 Chunk，符号与其它可见原文仍逐字可引用。"""
    document_ir = parse_package(content).document_ir
    whitespace_nodes = {
        node.node_id: node.text_payload.exact_text
        for node in document_ir.nodes
        if node.text_payload is not None
        and node.text_payload.exact_text
        and not node.text_payload.exact_text.strip()
        and node.kind in {NodeKind.PARAGRAPH, NodeKind.LIST_ITEM}
    }

    result = _chunk(document_ir)
    represented = {
        span.node_id
        for chunk in result.chunks
        for span in chunk.source_spans
        if span.node_id is not None
    }

    assert whitespace_nodes
    assert represented.isdisjoint(whitespace_nodes)
    assert all(chunk.citation_text.strip() for chunk in result.chunks)
    assert any(
        "。" in chunk.citation_text or "；" in chunk.citation_text
        for chunk in result.chunks
    )
    assert result.report.source_span_coverage == 1.0
    assert result.report.missing_source_chars == 0
    assert result.report.whitespace_only_citable_node_count == len(
        whitespace_nodes
    )
    assert result.report.whitespace_only_citable_char_count == sum(
        len(text) for text in whitespace_nodes.values()
    )


def test_empty_table_row_does_not_publish_separator_only_chunk() -> None:
    """全空表格行不伪造正文，含标点的相邻行仍保留真实行列绑定。"""
    table = """
<w:tbl><w:tblGrid><w:gridCol/><w:gridCol/></w:tblGrid>
  <w:tr>
    <w:tc><w:p><w:r><w:t xml:space="preserve">   </w:t></w:r></w:p></w:tc>
    <w:tc><w:p><w:r><w:t xml:space="preserve"> </w:t></w:r></w:p></w:tc>
  </w:tr>
  <w:tr>
    <w:tc><w:p/></w:tc>
    <w:tc><w:p><w:r><w:t>；</w:t></w:r></w:p></w:tc>
  </w:tr>
</w:tbl>
"""
    document_ir = parse_package(build_package(table)).document_ir
    result = _chunk(document_ir)
    rows = [
        node for node in document_ir.nodes if node.kind is NodeKind.TABLE_ROW
    ]
    represented = {
        span.node_id
        for chunk in result.chunks
        for span in chunk.source_spans
        if span.node_id is not None
    }
    nodes = {node.node_id: node for node in document_ir.nodes}
    represented_rows = {
        nodes[node_id].parent_node_id
        for node_id in represented
        if node_id in nodes
        and nodes[node_id].parent_node_id is not None
        and nodes[nodes[node_id].parent_node_id].kind is NodeKind.TABLE_CELL
    }
    represented_rows = {
        nodes[cell_id].parent_node_id
        for cell_id in represented_rows
        if cell_id is not None
    }

    assert len(rows) == 2
    assert rows[0].node_id not in represented_rows
    assert rows[1].node_id in represented_rows
    assert all(chunk.citation_text.strip() != "|" for chunk in result.chunks)
    assert any("；" in chunk.citation_text for chunk in result.chunks)
    assert result.report.table_row_count == 2
    assert result.report.represented_table_row_count == 1
    assert result.report.table_cell_count == 4
    assert result.report.represented_table_cell_count == 2


def test_missing_repeated_multilevel_and_long_headings_are_bounded() -> None:
    """无标题及同名、多级、长标题均保持来源覆盖和 hard max。"""
    plain_ir = parse_package(
        build_package("<w:p><w:r><w:t>无标题正文。</w:t></w:r></w:p>")
    ).document_ir
    plain = _chunk(plain_ir)
    assert all(not chunk.heading_path for chunk in plain.chunks)
    assert plain.report.source_span_coverage == 1.0

    long_heading = "复杂层级" * 24
    structured_ir = parse_package(
        build_package(
            '<w:p><w:pPr><w:outlineLvl w:val="0"/></w:pPr>'
            "<w:r><w:t>同名标题</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>第一段正文。</w:t></w:r></w:p>"
            '<w:p><w:pPr><w:outlineLvl w:val="1"/></w:pPr>'
            "<w:r><w:t>同名标题</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>第二段正文。</w:t></w:r></w:p>"
            '<w:p><w:pPr><w:outlineLvl w:val="0"/></w:pPr>'
            f"<w:r><w:t>{long_heading}</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>第三段正文。</w:t></w:r></w:p>"
        )
    ).document_ir
    structured = _chunk(structured_ir)
    headings = [
        node for node in structured_ir.nodes if node.kind is NodeKind.HEADING
    ]

    assert [node.text for node in headings] == [
        "同名标题",
        "同名标题",
        long_heading,
    ]
    assert [dict(node.metadata)["heading_level"] for node in headings] == [
        1,
        2,
        1,
    ]
    assert structured.report.source_span_coverage == 1.0
    assert structured.report.missing_source_chars == 0
    assert all(chunk.token_count <= 512 for chunk in structured.chunks)


def test_split_runs_preserve_same_semantics_with_nbsp_and_constraints() -> None:
    """不同 run 划分不改变含 NBSP、否定、单位和标识符的三视图。"""
    text = "设备 QX-91 不得超过 18\u00a0kPa。"
    single_ir = parse_package(
        build_package(f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>")
    ).document_ir
    split_ir = parse_package(
        build_package(
            "<w:p><w:r><w:t>设备 QX-91 </w:t></w:r>"
            "<w:r><w:t>不得超过 </w:t></w:r>"
            '<w:r><w:t xml:space="preserve">18\u00a0kPa。</w:t></w:r></w:p>'
        )
    ).document_ir
    single = _chunk(single_ir)
    split = _chunk(split_ir)

    assert single_ir.nodes[0].text_payload == split_ir.nodes[0].text_payload
    assert [
        (chunk.citation_text, chunk.embedding_text, chunk.lexical_text)
        for chunk in single.chunks
    ] == [
        (chunk.citation_text, chunk.embedding_text, chunk.lexical_text)
        for chunk in split.chunks
    ]
    assert text in single.chunks[0].citation_text


def test_list_soft_break_and_manual_number_keep_distinct_contracts() -> None:
    """自动列表内软换行保留映射，手工编号仍是普通正文。"""
    document_ir = parse_package(
        build_package(
            '<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/>'
            '<w:numId w:val="7"/></w:numPr></w:pPr>'
            "<w:r><w:t>第一步</w:t><w:br/><w:t>继续说明</w:t></w:r>"
            "</w:p>"
            "<w:p><w:r><w:t>A) 手工编号正文。</w:t></w:r></w:p>",
            numbering=_numbering(),
        )
    ).document_ir
    result = _chunk(document_ir)
    list_node = next(
        node for node in document_ir.nodes if node.kind is NodeKind.LIST_ITEM
    )
    manual_node = next(
        node
        for node in document_ir.nodes
        if node.kind is NodeKind.PARAGRAPH and node.text.startswith("A)")
    )
    original_spans = [
        span
        for chunk in result.chunks
        for span in chunk.source_spans
        if span.node_id == list_node.node_id
        and span.span_type is SourceSpanKind.ORIGINAL_TEXT
        and not span.is_repeated
    ]
    assert list_node.text_payload is not None
    assert list_node.text_payload.exact_text == "第一步\n继续说明"
    assert list_node.list_attributes is not None
    assert list_node.list_attributes.marker == "1. "
    assert manual_node.list_attributes is None
    assert (
        "".join(
            list_node.text_payload.exact_text[
                span.source_start_char : span.source_end_char
            ]
            for span in original_spans
        )
        == list_node.text_payload.exact_text
    )
    assert result.report.list_label_count == 1
    assert result.report.represented_list_label_count == 1
    assert result.report.source_span_coverage == 1.0


@pytest.mark.parametrize(
    ("delta", "expected_segments"),
    ((-1, 1), (0, 1), (1, 2)),
)
def test_actual_tokenizer_hard_boundary_minus_equal_plus_one(
    delta: int,
    expected_segments: int,
) -> None:
    """真实离线 tokenizer 在 hard max 两侧严格拆分且逐字前进。"""
    counter = DeterministicUtf8TokenCounter()
    hard_limit = 64
    fixed_prefix = "类型：正文\n\n"
    assert counter.count(fixed_prefix).count == 17
    text = "A" * (hard_limit - 17 + delta)
    anchor = SourceAnchor(
        part_uri="/word/document.xml",
        story_kind=StoryKind.BODY,
        structural_path=("body", "paragraph:0"),
        ordinal=0,
    )
    atom = AtomicUnit(
        unit_id=deterministic_id("node", "boundary", delta),
        role=ChunkRole.TEXT,
        parent_node_id=None,
        section_id="section-boundary",
        neighbor_group_id="group-boundary",
        heading_path=(),
        fragments=(
            SourceFragment(
                text=text,
                span_type=SourceSpanKind.ORIGINAL_TEXT,
                node_id=deterministic_id("node", "boundary-source", delta),
                source_anchor=anchor,
                source_start_char=0,
                source_end_char=len(text),
            ),
        ),
    )
    policy = ChunkingPolicy(
        target_tokens=hard_limit,
        hard_max_tokens=hard_limit,
        overlap_cap_tokens=0,
        min_tail_tokens=0,
        profile_hard_cap=hard_limit,
    )

    segments = split_atom(
        atom,
        document_title="",
        policy=policy,
        token_counter=counter,
    )
    rendered = tuple(render_fragments(item.fragments) for item in segments)

    assert len(segments) == expected_segments
    assert "".join(item.text for item in rendered) == text
    assert all(
        counter.count(embedding_text("", segment, item.text)).count
        <= hard_limit
        for segment, item in zip(segments, rendered, strict=True)
    )


def test_combining_characters_and_emoji_are_not_split_mid_grapheme() -> None:
    """组合字符与 emoji 在 UTF-8 硬预算拆分后仍保持字符簇边界。"""
    text = ("e\u0301🙂" * 24) + "完成。"
    node_id = deterministic_id("node", "unicode-boundary")
    anchor = SourceAnchor(
        part_uri="/word/document.xml",
        story_kind=StoryKind.BODY,
        structural_path=("body", "paragraph:unicode"),
        ordinal=0,
    )
    atom = AtomicUnit(
        unit_id=node_id,
        role=ChunkRole.TEXT,
        parent_node_id=node_id,
        section_id="section-unicode",
        neighbor_group_id="group-unicode",
        heading_path=(),
        fragments=(
            SourceFragment(
                text=text,
                span_type=SourceSpanKind.ORIGINAL_TEXT,
                node_id=node_id,
                source_anchor=anchor,
                source_start_char=0,
                source_end_char=len(text),
            ),
        ),
    )
    policy = ChunkingPolicy(
        target_tokens=54,
        hard_max_tokens=64,
        overlap_cap_tokens=0,
        min_tail_tokens=0,
        profile_hard_cap=64,
    )

    segments = split_atom(
        atom,
        document_title="",
        policy=policy,
        token_counter=DeterministicUtf8TokenCounter(),
    )
    texts = tuple(render_fragments(item.fragments).text for item in segments)

    assert len(texts) > 1
    assert "".join(texts) == text
    assert all(
        not value or unicodedata.combining(value[0]) == 0 for value in texts
    )


def test_long_table_row_never_emits_separator_only_segment() -> None:
    """长表格行跨块时，列分隔符不能独立成为无来源 Chunk 候选。"""
    first_node_id = deterministic_id("node", "long-table", "first")
    second_node_id = deterministic_id("node", "long-table", "second")
    first_anchor = SourceAnchor(
        part_uri="/word/document.xml",
        story_kind=StoryKind.BODY,
        structural_path=("body", "table:0", "row:0", "cell:0"),
        ordinal=0,
    )
    second_anchor = first_anchor.model_copy(
        update={
            "structural_path": (
                "body",
                "table:0",
                "row:0",
                "cell:1",
            ),
            "ordinal": 1,
        }
    )
    atom = AtomicUnit(
        unit_id=deterministic_id("node", "long-table", "row"),
        role=ChunkRole.TABLE,
        parent_node_id=deterministic_id("node", "long-table", "table"),
        section_id="section-long-table",
        neighbor_group_id="group-long-table",
        heading_path=(),
        fragments=(
            SourceFragment(
                text="A" * 495,
                span_type=SourceSpanKind.ORIGINAL_TEXT,
                node_id=first_node_id,
                source_anchor=first_anchor,
                source_start_char=0,
                source_end_char=495,
            ),
            separator_fragment(" | "),
            SourceFragment(
                text="B" * 495,
                span_type=SourceSpanKind.ORIGINAL_TEXT,
                node_id=second_node_id,
                source_anchor=second_anchor,
                source_start_char=0,
                source_end_char=495,
            ),
        ),
    )

    segments = split_atom(
        atom,
        document_title="",
        policy=ChunkingPolicy(),
        token_counter=DeterministicUtf8TokenCounter(),
    )

    assert len(segments) > 1
    assert all(
        any(
            fragment.span_type is not SourceSpanKind.SEPARATOR
            for fragment in segment.fragments
        )
        for segment in segments
    )
    assert "".join(
        fragment.text
        for segment in segments
        for fragment in segment.fragments
        if fragment.span_type is not SourceSpanKind.SEPARATOR
        and not fragment.is_repeated
    ) == ("A" * 495) + ("B" * 495)


def test_long_context_never_splits_derived_list_marker() -> None:
    """长表头上下文必须给完整派生编号预留真实 token 预算。"""
    preface_node_id = deterministic_id("node", "derived-marker-preface")
    node_id = deterministic_id("node", "derived-marker")
    preface_anchor = SourceAnchor(
        part_uri="/word/document.xml",
        story_kind=StoryKind.BODY,
        structural_path=("body", "table:0", "row:0", "cell:0"),
        ordinal=0,
    )
    anchor = preface_anchor.model_copy(
        update={
            "structural_path": (
                "body",
                "table:0",
                "row:0",
                "list:0",
            ),
            "ordinal": 1,
        }
    )
    marker = "2. "
    source = "逐项验证结构化结果。" * 80
    atom = AtomicUnit(
        unit_id=deterministic_id("node", "derived-marker-row"),
        role=ChunkRole.TABLE,
        parent_node_id=deterministic_id("node", "derived-marker-table"),
        section_id="section-derived-marker",
        neighbor_group_id="group-derived-marker",
        heading_path=("复杂表格",),
        fragments=(
            SourceFragment(
                text="甲" * 13,
                span_type=SourceSpanKind.ORIGINAL_TEXT,
                node_id=preface_node_id,
                source_anchor=preface_anchor,
                source_start_char=0,
                source_end_char=13,
            ),
            separator_fragment("\n"),
            SourceFragment(
                text=marker,
                span_type=SourceSpanKind.DERIVED_NUMBERING,
                node_id=node_id,
                source_anchor=anchor,
            ),
            SourceFragment(
                text=source,
                span_type=SourceSpanKind.ORIGINAL_TEXT,
                node_id=node_id,
                source_anchor=anchor,
                source_start_char=0,
                source_end_char=len(source),
            ),
        ),
        structural_context=("长表头" * 29) + "a",
    )

    segments = split_atom(
        atom,
        document_title="screenshot-input.docx",
        policy=ChunkingPolicy(),
        token_counter=DeterministicUtf8TokenCounter(),
    )
    derived = [
        fragment.text
        for segment in segments
        for fragment in segment.fragments
        if fragment.span_type is SourceSpanKind.DERIVED_NUMBERING
    ]

    assert derived == [marker]
    assert all(
        fragment.text == marker
        for segment in segments
        for fragment in segment.fragments
        if fragment.span_type is SourceSpanKind.DERIVED_NUMBERING
    )
    original_text = "".join(
        fragment.text
        for segment in segments
        for fragment in segment.fragments
        if fragment.span_type is SourceSpanKind.ORIGINAL_TEXT
        and not fragment.is_repeated
    )
    assert original_text == ("甲" * 13) + source


def test_long_list_label_report_counts_logical_nodes_not_segments() -> None:
    """长列表项跨块时逻辑编号只计一次，不能按片段膨胀。"""
    document_ir = parse_package(
        build_package(
            _list_item("逐项核验。" * 80, 7, 0),
            numbering=_numbering(),
        )
    ).document_ir
    policy = ChunkingPolicy(
        target_tokens=80,
        hard_max_tokens=120,
        overlap_cap_tokens=16,
        min_tail_tokens=10,
        profile_hard_cap=120,
    )

    result = _chunk(document_ir, policy=policy)
    labels = [
        span
        for chunk in result.chunks
        for span in chunk.source_spans
        if span.span_type is SourceSpanKind.DERIVED_NUMBERING
    ]

    assert len(result.chunks) > 1
    assert len(labels) == 1
    assert result.report.list_label_count == 1
    assert result.report.represented_list_label_count == 1


def test_chunk_json_roundtrip_revalidates_hash_and_source_contract() -> None:
    """Chunk JSON 回读重跑模型合同，损坏 hash 不能因序列化而逃逸。"""
    document_ir = parse_package(
        build_package("<w:p><w:r><w:t>合成回读内容</w:t></w:r></w:p>")
    ).document_ir
    chunk = _chunk(document_ir).chunks[0]

    assert Chunk.model_validate_json(chunk.model_dump_json()) == chunk
    damaged = chunk.model_dump(mode="json")
    damaged["content_sha256"] = hashlib.sha256(b"damaged").hexdigest()
    with pytest.raises(ValidationError, match="content_sha256"):
        Chunk.model_validate(damaged)


def test_validator_rejects_duplicate_chunk_ids() -> None:
    """最终 validator 拒绝重复稳定 ID，不能只在报告里记数。"""
    document_ir = parse_package(
        build_package("<w:p><w:r><w:t>独立重复块样例</w:t></w:r></w:p>")
    ).document_ir
    result = _chunk(document_ir)
    chunk = result.chunks[0]

    with pytest.raises(ValueError, match="稳定 chunk ID 禁止重复"):
        validate_chunks(
            (chunk, chunk),
            document_ir,
            ChunkingPolicy(),
            DeterministicUtf8TokenCounter(),
        )
