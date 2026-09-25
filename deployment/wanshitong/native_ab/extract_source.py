#!/usr/bin/env python3
"""从冻结 DOCX 提取带定位和摘要的标注候选，不写入仓库。"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path

from docx import Document


def main() -> None:
    """逐份校验原件，再提取正文段落与表格单元格。"""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--lock", type=Path, required=True)
    parser.add_argument("--corpus", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    lock = json.loads(args.lock.read_text(encoding="utf-8"))
    descriptor = os.open(
        args.output, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600
    )
    counts: dict[str, int] = {}
    with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
        for entry in lock["documents"]:
            path = args.corpus / entry["source_relative_path"]
            raw = path.read_bytes()
            if hashlib.sha256(raw).hexdigest() != entry["sha256"]:
                raise ValueError(f"原件 SHA 不符：{entry['control_id']}")
            document = Document(path)
            paragraphs = (
                (f"p:{index}", paragraph.text)
                for index, paragraph in enumerate(document.paragraphs)
            )
            cells = (
                (f"tbl:{table_index}/tr:{row_index}/tc:{cell_index}", cell.text)
                for table_index, table in enumerate(document.tables)
                for row_index, row in enumerate(table.rows)
                for cell_index, cell in enumerate(row.cells)
            )
            count = 0
            for locator, value in (*paragraphs, *cells):
                text = " ".join(value.split())
                if not text:
                    continue
                record = {
                    "document_id": entry["control_id"],
                    "document_sha256": entry["sha256"],
                    "locator": locator,
                    "excerpt_sha256": hashlib.sha256(text.encode()).hexdigest(),
                    "text": text,
                }
                stream.write(json.dumps(record, ensure_ascii=False) + "\n")
                count += 1
            counts[entry["control_id"]] = count
    print(json.dumps(counts, ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    main()
